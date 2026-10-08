"""Importación desde Excel: listado maestro (hojas LM-AAAA) y precios históricos (hojas L1..L11)."""
import io
import re
import unicodedata

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Articulo, Lote, PrecioHistorico

UNIDADES_SECOP = {
    "UN", "UND", "KG", "G", "L", "LT", "ML", "CC", "GAR", "SC", "PAC", "ROL", "CX", "CXC", "PAR", "BDN",
    "BOB", "LB", "BTO", "CON", "M", "M2", "M3", "GL", "GAL", "KL", "FR", "CJ", "LAM", "MAC", "DOC", "JGO",
}
SINONIMOS_UNIDAD = {"UNIDAD": "UN", "KL": "KG", "ROLL": "ROL", "LITRO": "L"}


def normalizar(txt: str) -> str:
    txt = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", txt).strip()


def limpiar_codigo(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", str(v).strip())


def limpiar_unidad(v) -> str:
    u = str(v or "").strip().upper()
    return SINONIMOS_UNIDAD.get(u, u)


def abrir_libro(contenido: bytes):
    try:
        return load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "El archivo no es un Excel (.xlsx) válido")


def hojas_listado(wb) -> list[str]:
    return [ws.title for ws in wb.worksheets if _fila_encabezado(ws) is not None]


def _fila_encabezado(ws, max_filas: int = 15):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_filas, values_only=True), 1):
        textos = [normalizar(c) for c in row if c is not None]
        if any("unspsc" in t for t in textos) and any(t.startswith("producto") for t in textos):
            return i, [normalizar(c) if c is not None else "" for c in row]
    return None


def _col(enc: list[str], *claves: str) -> int | None:
    for i, t in enumerate(enc):
        if any(t.startswith(k) or k in t for k in claves):
            return i
    return None


def leer_listado(wb, hoja: str) -> list[dict]:
    """Lee una hoja con columnas Lote | Código UNSPSC | Producto | Unidad | Descripción."""
    if hoja not in wb.sheetnames:
        raise HTTPException(422, f"La hoja '{hoja}' no existe")
    ws = wb[hoja]
    found = _fila_encabezado(ws)
    if not found:
        raise HTTPException(422, f"La hoja '{hoja}' no tiene encabezados de listado maestro (UNSPSC, Producto)")
    fila, enc = found
    c_lote = _col(enc, "lote")
    c_cod = _col(enc, "unspsc")
    c_prod = _col(enc, "producto")
    c_und = _col(enc, "unidad")
    c_desc = _col(enc, "descripcion")
    filas = []
    for n, row in enumerate(ws.iter_rows(min_row=fila + 1, values_only=True), fila + 1):
        if c_prod is None or c_prod >= len(row) or not row[c_prod]:
            continue
        filas.append({
            "fila": n,
            "lote": str(row[c_lote]).strip() if c_lote is not None and row[c_lote] is not None else "",
            "codigo_unspsc": limpiar_codigo(row[c_cod]) if c_cod is not None else "",
            "nombre": re.sub(r"\s+", " ", str(row[c_prod])).strip(),
            "unidad": limpiar_unidad(row[c_und]) if c_und is not None else "",
            "descripcion": str(row[c_desc] or "").strip() if c_desc is not None and c_desc < len(row) else "",
        })
    return filas


def previsualizar_listado(db: Session, filas: list[dict], desactivar_ausentes: bool, centro_id: int) -> dict:
    """Compara el archivo con el listado maestro del centro: los lotes se buscan sólo entre los del centro."""
    lotes = {l.abreviatura.lower(): l for l in db.scalars(select(Lote).where(Lote.centro_id == centro_id))}
    lotes.update({str(l.numero): l for l in lotes.copy().values()})
    existentes = {(a.lote_id, normalizar(a.nombre)): a
                  for a in db.scalars(select(Articulo).join(Lote).where(Lote.centro_id == centro_id))}

    nuevos, actualizados, sin_cambios, errores, advertencias = [], [], [], [], []
    vistos: set = set()
    for f in filas:
        lote = lotes.get(f["lote"].lower())
        errs = []
        if not lote:
            errs.append(f"Lote '{f['lote']}' no existe")
        if not re.fullmatch(r"\d{8}", f["codigo_unspsc"]):
            errs.append(f"Código UNSPSC inválido '{f['codigo_unspsc']}' (deben ser 8 dígitos)")
        if not f["unidad"] or f["unidad"] == "0":
            errs.append("Unidad de medida vacía")
        if errs:
            errores.append({**f, "errores": errs})
            continue
        if f["unidad"] not in UNIDADES_SECOP:
            advertencias.append({**f, "mensaje": f"Unidad '{f['unidad']}' no está en la lista habitual de SECOP II"})
        clave = (lote.id, normalizar(f["nombre"]))
        if clave in vistos:
            advertencias.append({**f, "mensaje": "Producto duplicado en el archivo (se toma la primera aparición)"})
            continue
        vistos.add(clave)
        f = {**f, "lote_id": lote.id, "lote_nombre": lote.nombre}
        a = existentes.get(clave)
        if not a:
            nuevos.append(f)
        else:
            cambios = [k for k in ("codigo_unspsc", "unidad", "descripcion") if (getattr(a, k) or "") != f[k]]
            if not a.activo:
                cambios.append("activo")
            (actualizados if cambios else sin_cambios).append({**f, "id": a.id, "cambios": cambios})

    ausentes = []
    if desactivar_ausentes:
        lotes_archivo = {lid for lid, _ in vistos}
        ausentes = [
            {"id": a.id, "nombre": a.nombre, "lote_id": a.lote_id}
            for (lid, nom), a in existentes.items()
            if lid in lotes_archivo and (lid, nom) not in vistos and a.activo
        ]
    return {
        "resumen": {"leidas": len(filas), "nuevos": len(nuevos), "actualizados": len(actualizados),
                    "sin_cambios": len(sin_cambios), "errores": len(errores),
                    "advertencias": len(advertencias), "a_desactivar": len(ausentes)},
        "nuevos": nuevos, "actualizados": actualizados, "errores": errores,
        "advertencias": advertencias, "a_desactivar": ausentes,
    }


def aplicar_listado(db: Session, prev: dict) -> dict:
    for f in prev["nuevos"]:
        db.add(Articulo(lote_id=f["lote_id"], codigo_unspsc=f["codigo_unspsc"], nombre=f["nombre"][:300],
                        unidad=f["unidad"][:10], descripcion=f["descripcion"], activo=True))
    for f in prev["actualizados"]:
        a = db.get(Articulo, f["id"])
        a.codigo_unspsc, a.unidad, a.descripcion, a.activo = f["codigo_unspsc"], f["unidad"][:10], f["descripcion"], True
    for f in prev["a_desactivar"]:
        db.get(Articulo, f["id"]).activo = False
    return prev["resumen"]


def importar_historicos_estudio(db: Session, wb, anio: int, fuente: str, centro_id: int) -> int:
    """Toma el 'Precio Unitario Estimado IVA Incluido' de las hojas L1..L11 de un estudio
    de mercados ya firmado y lo registra como precio histórico del año indicado (artículos del centro)."""
    articulos = {(a.lote_id, normalizar(a.nombre)): a.id
                 for a in db.scalars(select(Articulo).join(Lote).where(Lote.centro_id == centro_id))}
    lotes = {l.abreviatura.lower(): l.id for l in db.scalars(select(Lote).where(Lote.centro_id == centro_id))}
    ya = {(h.articulo_id, h.anio) for h in db.scalars(select(PrecioHistorico).where(PrecioHistorico.anio == anio))}
    total = 0
    for ws in wb.worksheets:
        if not re.fullmatch(r"L\d{1,2}", ws.title):
            continue
        found = _fila_encabezado(ws)
        if not found:
            continue
        fila, enc = found
        c_prod, c_lote = _col(enc, "producto"), _col(enc, "abrev")
        c_precio = _col(enc, "precio unitario estimado")
        if None in (c_prod, c_lote, c_precio):
            continue
        for row in ws.iter_rows(min_row=fila + 1, values_only=True):
            if not row[c_prod] or row[c_lote] is None:
                continue
            try:
                precio = float(row[c_precio])
            except (TypeError, ValueError):
                continue
            aid = articulos.get((lotes.get(str(row[c_lote]).strip().lower()), normalizar(row[c_prod])))
            if aid and precio > 0 and (aid, anio) not in ya:
                db.add(PrecioHistorico(articulo_id=aid, anio=anio, precio=round(precio, 2), fuente=fuente))
                ya.add((aid, anio))
                total += 1
    return total


def leer_columnas_id_valor(contenido: bytes, columna_valor: str) -> dict[int, float]:
    """Lee una plantilla generada por el sistema (columna ID + columna de valor)."""
    wb = abrir_libro(contenido)
    ws = wb.worksheets[0]
    enc_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        enc = [normalizar(c) if c is not None else "" for c in row]
        if "id" in enc and any(t.startswith(columna_valor) for t in enc):
            enc_idx = (i, enc.index("id"), next(j for j, t in enumerate(enc) if t.startswith(columna_valor)))
            break
    if not enc_idx:
        raise HTTPException(422, f"La plantilla debe tener las columnas 'ID' y '{columna_valor.title()}'")
    fila, c_id, c_val = enc_idx
    out: dict[int, float] = {}
    for row in ws.iter_rows(min_row=fila + 1, values_only=True):
        if row[c_id] is None or row[c_val] in (None, ""):
            continue
        try:
            out[int(row[c_id])] = float(str(row[c_val]).replace("$", "").replace(".", "").replace(",", ".")
                                        if isinstance(row[c_val], str) else row[c_val])
        except (TypeError, ValueError):
            raise HTTPException(422, f"Valor no numérico para el ID {row[c_id]}: '{row[c_val]}'")
    return out
