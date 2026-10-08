"""Generación de archivos Excel con formato institucional."""
import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERDE, AZUL, GRIS = "39A909", "1E3A5F", "EDF2F7"
_borde = Border(*(Side(style="thin", color="C8D0DA"),) * 4)
_enc_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
_enc_fill = PatternFill("solid", fgColor=AZUL)
MONEDA = '"$"#,##0'
LOGO = Path(__file__).resolve().parent.parent / "static" / "logo-sena.png"
PORC = "0.00%"


def _titulo(ws, lineas: list[str], ancho: int, logo: bool = True):
    for i, txt in enumerate(lineas, 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ancho)
        c = ws.cell(i, 1, txt)
        c.font = Font(bold=True, size=12 if i == 1 else 10, color=VERDE if i == 1 else "1A202C")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    if logo and LOGO.exists():
        img = ImagenExcel(str(LOGO))
        img.width, img.height = round(56 * img.width / img.height), 56
        ws.add_image(img, "A1")
        for i in range(1, len(lineas) + 1):  # filas suficientes para el logo
            ws.row_dimensions[i].height = max(ws.row_dimensions[i].height or 15, 64 / len(lineas))
    return len(lineas) + 2


def _tabla(ws, fila: int, encabezados: list[tuple[str, int]], filas: list[list], formatos: dict[int, str] | None = None):
    formatos = formatos or {}
    for j, (txt, ancho) in enumerate(encabezados, 1):
        c = ws.cell(fila, j, txt)
        c.font, c.fill, c.border = _enc_font, _enc_fill, _borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila].height = 32
    for i, valores in enumerate(filas, fila + 1):
        for j, v in enumerate(valores, 1):
            c = ws.cell(i, j, v)
            c.border = _borde
            c.alignment = Alignment(vertical="top", wrap_text=j in formatos.get("wrap", ()))  # type: ignore[arg-type]
            if j in formatos:
                c.number_format = formatos[j]
            if (i - fila) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GRIS)
    ws.freeze_panes = ws.cell(fila + 1, 1)
    if filas:
        ws.auto_filter.ref = f"A{fila}:{get_column_letter(len(encabezados))}{fila + len(filas)}"
    return fila + len(filas) + 1


def guardar(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def encabezado(centros) -> list[str]:
    """Encabezado institucional con el centro de formación (o genérico si el archivo abarca varios centros)."""
    centros = {c.id: c for c in centros}
    if len(centros) == 1:
        c = next(iter(centros.values()))
        linea = f"{c.nombre} - Regional {c.regional}" if c.regional else c.nombre
    else:
        linea = "Centros de formación"
    return ["SERVICIO NACIONAL DE APRENDIZAJE - SENA", linea]


def listado_maestro(articulos: list, titulo: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Listado maestro"
    f = _titulo(ws, encabezado(a.lote.centro for a in articulos) + [titulo, f"Generado el {date.today():%d/%m/%Y}"], 7)
    filas = [[a.lote.numero, a.lote.abreviatura, a.codigo_unspsc, a.nombre, a.unidad, a.descripcion,
              "Activo" if a.activo else "Inactivo"] for a in articulos]
    _tabla(ws, f, [("N.", 5), ("Lote", 8), ("Código UNSPSC", 12), ("Producto", 45), ("Unidad SECOP", 10),
                   ("Descripción", 80), ("Estado", 9)], filas, {"wrap": (4, 6)})
    return guardar(wb)


def plantilla(titulo: str, articulos: list, columna: str, valores: dict[int, float] | None = None,
              instruccion: str = "") -> bytes:
    """Plantilla de diligenciamiento: el usuario sólo llena la última columna."""
    valores = valores or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    f = _titulo(ws, [titulo, instruccion or f"Diligencie únicamente la columna '{columna}'. No modifique la columna ID."], 6,
                logo=False)
    filas = [[a.id, a.codigo_unspsc, a.nombre, a.unidad, a.descripcion[:250], valores.get(a.id)] for a in articulos]
    _tabla(ws, f, [("ID", 7), ("Código UNSPSC", 12), ("Producto", 45), ("Unidad", 8), ("Descripción", 60),
                   (columna, 18)], filas, {6: MONEDA if "precio" in columna.lower() else "#,##0.##", "wrap": (3,)})
    amarillo = PatternFill("solid", fgColor="FFF4C2")
    for i in range(f + 1, f + 1 + len(filas)):
        ws.cell(i, 6).fill = amarillo
    return guardar(wb)


def analisis_precios(vigencia, lote, articulos: list, resultados: dict, etiquetas: list[tuple[str, str]]) -> bytes:
    """Réplica del FORMATO ANÁLISIS DE PRECIOS Y CONSOLIDACIÓN (hojas L1..L11)."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"L{lote.numero}"
    anios = [vigencia.anio - 1, vigencia.anio - 2, vigencia.anio - 3]
    enc = ([("Lote", 5), ("Abrev", 7), ("Código UNSPSC", 11), ("Producto", 38), ("Unidad SECOP II", 8)]
           + [(f"Precio {e}", 12) for e, _ in etiquetas]
           + [(f"Histórico {a}*", 12) for a in anios]
           + [("Precio de experto", 12), ("Número de cotizaciones", 10), ("Precio promedio IVA incluido", 13),
              ("Dispersión inicial", 10), ("Análisis de los precios cotizados", 50),
              ("Precio promedio IVA incluido (final)", 13), ("Dispersión final", 10),
              ("Precio unitario estimado IVA incluido", 14), ("Descripción ficha técnica", 70)])
    f = _titulo(ws, encabezado([lote.centro]) + [
        "FORMATO ANÁLISIS DE PRECIOS Y CONSOLIDACIÓN DE BIENES DE CARACTERÍSTICAS TÉCNICAS UNIFORMES",
        f"LOTE {lote.numero} - {lote.nombre.upper()} ({lote.abreviatura}) · VIGENCIA {vigencia.codigo}",
    ], len(enc))
    filas = []
    for a in articulos:
        r = resultados[a.id]
        precios = {c["etiqueta"]: c["precio"] for c in r.cotizaciones}
        hist = {h["anio"]: h["indexado"] for h in r.historicos}
        n = r.n_validas if r.metodo == "COTIZACIONES" else len(r.historicos) if r.metodo == "HISTORICOS" else 0
        filas.append([lote.numero, lote.abreviatura, a.codigo_unspsc, a.nombre, a.unidad]
                     + [precios.get(e) for e, _ in etiquetas]
                     + [hist.get(y) for y in anios]
                     + [r.precio_experto, n, r.precio_promedio_inicial, r.dispersion_inicial, r.justificacion,
                        r.precio_estimado, r.dispersion_final, r.precio_estimado, a.descripcion])
    nc = len(etiquetas)
    fmt: dict = {i: MONEDA for i in range(6, 6 + nc + 4)}
    fmt.update({6 + nc + 5: MONEDA, 6 + nc + 6: PORC, 6 + nc + 8: MONEDA, 6 + nc + 9: PORC, 6 + nc + 10: MONEDA})
    fmt["wrap"] = (4, 6 + nc + 7, 6 + nc + 11)
    fin = _tabla(ws, f, enc, filas, fmt)
    ws.cell(fin + 1, 1, "Listado de empresas cotizantes.").font = Font(bold=True)
    for i, (e, prov) in enumerate(etiquetas, fin + 2):
        ws.cell(i, 1, f"{e} = {prov}")
    nota = fin + 3 + len(etiquetas)
    ws.cell(nota, 1, "* Los precios históricos fueron indexados por el Porcentaje de Ajuste Año Gravable (PAAG): "
                     "VF = VP*(1+IPC1)*(1+IPC2)*...*(1+IPCn), con IPC más los puntos adicionales configurados.")
    ws.cell(nota + 2, 1, "De conformidad con lo previsto en los artículos 2.2.1.1.1.6.1 y 2.2.1.1.2.1.1 del Decreto "
                         "1082 de 2015 y el Manual de Contratación Administrativa - SENA (GCON-M-001), se deja "
                         "constancia de la realización del estudio de mercado como parte integral del análisis del sector.")
    ws.cell(nota + 6, 2, "_______________________________")
    ws.cell(nota + 7, 2, "Elaboró")
    ws.cell(nota + 6, 5, "_______________________________")
    ws.cell(nota + 7, 5, "Revisó")
    return guardar(wb)


def consolidado(vigencia, solicitudes: list, precios: dict[int, float | None]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    f = _titulo(ws, encabezado(s.lote.centro for s in solicitudes) + [f"CONSOLIDADO DE SOLICITUDES - VIGENCIA {vigencia.codigo}",
                                       f"Generado el {date.today():%d/%m/%Y}"], 7)
    resumen = []
    for s in solicitudes:
        total = sum((i.precio_aprobado or precios.get(i.articulo_id) or 0) * i.cantidad for i in s.items)
        resumen.append([s.id, s.lote.centro.nombre, f"{s.lote.numero}. {s.lote.nombre}", s.lider.nombre,
                        s.estado_post.nombre if s.estado_post else s.estado, len(s.items), total,
                        s.fecha_envio.strftime("%d/%m/%Y") if s.fecha_envio else ""])
    fin = _tabla(ws, f, [("ID", 6), ("Centro de formación", 34), ("Lote", 34), ("Líder", 28), ("Estado", 14),
                         ("Artículos", 10), ("Valor estimado", 18), ("Fecha envío", 12)], resumen, {7: MONEDA})
    ws.cell(fin, 6, "TOTAL").font = Font(bold=True)
    c = ws.cell(fin, 7, sum(r[6] for r in resumen))
    c.number_format, c.font = MONEDA, Font(bold=True)

    det = wb.create_sheet("Detalle")
    filas = []
    for s in solicitudes:
        for i in s.items:
            p = i.precio_aprobado or precios.get(i.articulo_id)
            filas.append([s.lote.centro.codigo, s.lote.abreviatura, i.articulo.codigo_unspsc, i.articulo.nombre, i.articulo.unidad,
                          i.cantidad, p, (p or 0) * i.cantidad, s.estado_post.nombre if s.estado_post else s.estado])
    _tabla(det, 1, [("Centro", 10), ("Lote", 8), ("Código UNSPSC", 12), ("Producto", 50), ("Unidad", 8),
                    ("Cantidad", 10), ("Precio unitario", 14), ("Subtotal", 16), ("Estado", 14)], filas,
           {7: MONEDA, 8: MONEDA, "wrap": (4,)})
    return guardar(wb)
