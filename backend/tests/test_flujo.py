import io
import json

from openpyxl import Workbook, load_workbook

from app.services.pricing import cv, factor_indexacion
from .conftest import PDF, login


def _ids(client, h):
    arts = client.get("/api/articulos", headers=h).json()["items"]
    return {a["nombre"]: a["id"] for a in arts}


def _provs(client, h):
    return [p["id"] for p in client.get("/api/proveedores", headers=h).json()]


def _cotizar(client, h, prov, precios, fecha="2026-09-10"):
    return client.post("/api/cotizaciones", headers=h, data={"proveedor_id": prov, "fecha": fecha,
                                                             "precios": json.dumps(precios)},
                       files={"pdf": ("c.pdf", PDF, "application/pdf")})


# --- Reglas de cálculo (replican el Excel) --------------------------------
def test_dispersion_igual_que_excel():
    # Fila "Aceite Agricola" de la hoja L1: 29650, 33761, 37137 -> 0.11187...
    assert abs(cv([29650, 33761, 37137]) - 0.11187218346674402) < 1e-12
    assert cv([5]) == 0.0 and cv([]) is None


def test_indexacion_paag():
    class I:
        def __init__(self, ipc):
            self.ipc, self.puntos_adicionales = ipc, 0.02
    f = factor_indexacion({2023: I(0.0928), 2024: I(0.052), 2025: I(0.051)}, 2023, 2026, 0.05, 0.02)
    assert abs(f - 1.1128 * 1.072 * 1.071) < 1e-9


# --- Autenticación -------------------------------------------------------
def test_bloqueo_tras_tres_intentos(client):
    for _ in range(3):
        assert client.post("/api/auth/login", json={"email": "admin@sena.edu.co", "password": "x"}).status_code == 401
    assert client.post("/api/auth/login", json={"email": "admin@sena.edu.co", "password": "Admin123!"}).status_code == 423


def test_roles(client, lider):
    assert client.get("/api/usuarios", headers=lider).status_code == 403
    assert client.get("/api/usuarios").status_code == 401


def test_lider_requiere_lote(client, admin):
    r = client.post("/api/usuarios", headers=admin, json={"nombre": "X", "email": "x@sena.edu.co", "rol": "LIDER",
                                                          "password": "Clave123!"})
    assert r.status_code == 422


# --- Vigencias -----------------------------------------------------------
def test_una_sola_vigencia_abierta(client, manager):
    v = client.post("/api/vigencias", headers=manager, json={"codigo": "2028", "anio": 2028, "fecha_inicio": "2027-09-01",
                                                             "fecha_cierre": "2027-11-30"}).json()
    assert client.post(f"/api/vigencias/{v['id']}/abrir", headers=manager).status_code == 409


# --- Cotizaciones --------------------------------------------------------
def test_cotizacion_reglas(client, lider):
    ids, provs = _ids(client, lider), _provs(client, lider)
    aceite = ids["Aceite agrícola"]
    # PDF obligatorio y válido
    r = client.post("/api/cotizaciones", headers=lider, data={"proveedor_id": provs[0], "fecha": "2026-09-10",
                                                               "precios": json.dumps({aceite: 1000})},
                    files={"pdf": ("c.pdf", b"no es pdf", "application/pdf")})
    assert r.status_code == 422
    assert _cotizar(client, lider, provs[0], {aceite: 29650}).status_code == 201
    # mismo proveedor dos veces en el lote/vigencia
    assert _cotizar(client, lider, provs[0], {aceite: 1}).status_code == 409
    # fecha futura
    assert _cotizar(client, lider, provs[1], {aceite: 1}, fecha="2099-01-01").status_code == 422
    assert _cotizar(client, lider, provs[1], {aceite: 33761}).status_code == 201
    assert _cotizar(client, lider, provs[2], {aceite: 37137}).status_code == 201
    # máximo 3 cotizaciones por artículo
    r = _cotizar(client, lider, provs[3], {aceite: 30000})
    assert r.status_code == 409 and "ya tienen 3" in r.json()["detail"]


def test_carga_precios_por_plantilla_excel(client, lider):
    vig = client.get("/api/vigencias/activa", headers=lider).json()
    r = client.get(f"/api/cotizaciones/plantilla?vigencia_id={vig['id']}", headers=lider)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    enc = next(i for i in range(1, 10) if ws.cell(i, 1).value == "ID")
    ws.cell(enc + 1, 6, 15000)
    ws.cell(enc + 2, 6, 2500)
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/cotizaciones", headers=lider, data={"proveedor_id": _provs(client, lider)[0], "fecha": "2026-09-01"},
                    files={"pdf": ("c.pdf", PDF, "application/pdf"),
                           "archivo_precios": ("p.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 201, r.text
    assert r.json()["n_items"] == 2


# --- Análisis de precios --------------------------------------------------
def test_analisis_y_exclusion(client, lider, admin):
    ids, provs = _ids(client, lider), _provs(client, lider)
    az = ids["Azufre mineral"]
    _cotizar(client, lider, provs[0], {az: 7945})
    _cotizar(client, lider, provs[1], {az: 8100})
    _cotizar(client, lider, provs[2], {az: 21321})
    vig = client.get("/api/vigencias/activa", headers=admin).json()["id"]
    lote = client.get("/api/auth/me", headers=lider).json()["lote_id"]
    item = next(i for i in client.get(f"/api/analisis?vigencia_id={vig}&lote_id={lote}", headers=admin).json()["items"]
                if i["articulo_id"] == az)
    assert item["metodo"] == "COTIZACIONES" and item["soporte"] == "COMPLETO"
    assert item["semaforo"] == "rojo" and item["atipico_sugerido"]
    sugerido = next(c for c in item["cotizaciones"] if c["item_id"] == item["atipico_sugerido"])
    assert sugerido["precio"] == 21321
    # exclusión exige motivo
    assert client.post(f"/api/cotizaciones/items/{sugerido['item_id']}/exclusion", headers=admin,
                       json={"excluido": True}).status_code == 422
    assert client.post(f"/api/cotizaciones/items/{sugerido['item_id']}/exclusion", headers=admin,
                       json={"excluido": True, "motivo": "su valor notoriamente alto"}).status_code == 200
    item = next(i for i in client.get(f"/api/analisis?vigencia_id={vig}&lote_id={lote}", headers=admin).json()["items"]
                if i["articulo_id"] == az)
    assert item["precio_estimado"] == round((7945 + 8100) / 2, 2)
    assert item["justificacion"].startswith("Se excluye P3 por su valor notoriamente alto")
    # el líder no puede excluir
    assert client.post(f"/api/cotizaciones/items/{sugerido['item_id']}/exclusion", headers=lider,
                       json={"excluido": False}).status_code == 403
    r = client.get(f"/api/analisis/exportar?vigencia_id={vig}&lote_id={lote}", headers=admin)
    assert r.status_code == 200 and load_workbook(io.BytesIO(r.content)).active.title == "L1"


def test_historicos_cuando_faltan_cotizaciones(client, admin):
    from app.db import SessionLocal
    from app.models import Articulo, PrecioHistorico
    with SessionLocal() as db:
        a = db.query(Articulo).filter_by(nombre="Bacteria Bacillus subtilis").one()
        db.add(PrecioHistorico(articulo_id=a.id, anio=2026, precio=50000))
        db.commit()
        aid, lote = a.id, a.lote_id
    vig = client.get("/api/vigencias/activa", headers=admin).json()["id"]
    item = next(i for i in client.get(f"/api/analisis?vigencia_id={vig}&lote_id={lote}", headers=admin).json()["items"]
                if i["articulo_id"] == aid)
    assert item["metodo"] == "HISTORICOS"
    assert item["precio_estimado"] == round(50000 * (1 + 0.051 + 0.02), 2)  # 2026 sin índice -> parámetro por defecto


# --- Solicitud por lote: flujo completo -----------------------------------
def test_flujo_solicitud_completo(client, lider, admin):
    ids, provs = _ids(client, lider), _provs(client, lider)
    aceite, adaptador = ids["Aceite agrícola"], ids["Adaptador dentado 1/2"]
    _cotizar(client, lider, provs[0], {aceite: 29650})
    _cotizar(client, lider, provs[1], {aceite: 33761})

    mia = client.get("/api/solicitudes/mia", headers=lider).json()
    sid = mia["solicitud"]["id"]
    assert mia["solicitud"]["estado"] == "BORRADOR"
    # autoguardado por lotes
    r = client.put(f"/api/solicitudes/{sid}/items", headers=lider,
                   json=[{"articulo_id": aceite, "cantidad": 10}, {"articulo_id": adaptador, "cantidad": 5}])
    assert r.json()["agregados"] == 2
    # no se puede enviar: el adaptador no tiene precio
    r = client.post(f"/api/solicitudes/{sid}/enviar", headers=lider)
    assert r.status_code == 422 and any("Adaptador" in e for e in r.json()["detail"]["errores"])
    client.put(f"/api/solicitudes/{sid}/items", headers=lider, json=[{"articulo_id": adaptador, "cantidad": 0}])
    assert client.post(f"/api/solicitudes/{sid}/enviar", headers=lider).status_code == 200
    # bloqueada para edición y para nuevas cotizaciones del lote
    assert client.put(f"/api/solicitudes/{sid}/items", headers=lider,
                      json=[{"articulo_id": aceite, "cantidad": 1}]).status_code == 409
    assert _cotizar(client, lider, provs[2], {aceite: 30000}).status_code == 409

    assert client.post(f"/api/solicitudes/{sid}/revisar", headers=admin).status_code == 200
    # devolver exige observaciones
    assert client.post(f"/api/solicitudes/{sid}/decidir", headers=admin, json={"accion": "DEVOLVER"}).status_code == 422
    client.post(f"/api/solicitudes/{sid}/decidir", headers=admin,
                json={"accion": "DEVOLVER", "observaciones": "Ajustar cantidad"})
    client.put(f"/api/solicitudes/{sid}/items", headers=lider, json=[{"articulo_id": aceite, "cantidad": 8}])
    s = client.post(f"/api/solicitudes/{sid}/enviar", headers=lider).json()
    assert s["reenvios"] == 1

    d = client.post(f"/api/solicitudes/{sid}/decidir", headers=admin, json={"accion": "APROBAR"}).json()
    assert d["estado"] == "APROBADA"
    assert d["items"][0]["precio_congelado"] and d["items"][0]["precio_unitario"] == (29650 + 33761) / 2
    assert d["total"] == round(8 * (29650 + 33761) / 2, 2)

    estados = client.get("/api/estados-post", headers=admin).json()
    assert client.post(f"/api/solicitudes/{sid}/avanzar", headers=admin,
                       json={"estado_post_id": estados[1]["id"], "observaciones": ""}).status_code == 422
    client.post(f"/api/solicitudes/{sid}/avanzar", headers=admin,
                json={"estado_post_id": estados[1]["id"], "observaciones": "Despachado"})
    # sólo avanza
    r = client.post(f"/api/solicitudes/{sid}/avanzar", headers=admin,
                    json={"estado_post_id": estados[0]["id"], "observaciones": "x"})
    assert r.status_code == 409
    det = client.get(f"/api/solicitudes/{sid}", headers=lider).json()
    assert det["estado_visible"] == "En Camino"
    assert [h["estado_nuevo"] for h in det["historial"]] == [
        "BORRADOR", "ENVIADA", "EN_REVISION", "DEVUELTA", "ENVIADA", "APROBADA", "En Camino"]
    notifs = client.get("/api/notificaciones", headers=lider).json()
    assert notifs["no_leidas"] >= 3
    r = client.get(f"/api/solicitudes/consolidado/exportar?vigencia_id={det['vigencia_id']}", headers=admin)
    assert r.status_code == 200


def test_importar_cantidades_y_clonar(client, lider, admin, manager):
    ids, provs = _ids(client, lider), _provs(client, lider)
    aceite, azufre = ids["Aceite agrícola"], ids["Azufre mineral"]
    _cotizar(client, lider, provs[0], {aceite: 1000, azufre: 2000})
    sid = client.get("/api/solicitudes/mia", headers=lider).json()["solicitud"]["id"]
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Producto", "Cantidad"])
    ws.append([aceite, "Aceite", 4])
    ws.append([azufre, "Azufre", 12])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/solicitudes/{sid}/importar", headers=lider, files={"archivo": ("s.xlsx", buf.getvalue())})
    assert r.status_code == 200 and r.json()["n_items"] == 2
    client.post(f"/api/solicitudes/{sid}/enviar", headers=lider)
    client.post(f"/api/solicitudes/{sid}/decidir", headers=admin, json={"accion": "APROBAR"})

    # Nueva vigencia: cerrar 2027, abrir 2028 y clonar
    vigs = client.get("/api/vigencias", headers=manager).json()
    actual = next(v for v in vigs if v["estado"] == "ABIERTA")
    client.post(f"/api/vigencias/{actual['id']}/cerrar", headers=manager)
    nv = client.post("/api/vigencias", headers=manager, json={"codigo": "2028", "anio": 2028,
                                                              "fecha_inicio": "2027-09-01", "fecha_cierre": "2027-11-30"}).json()
    client.post(f"/api/vigencias/{nv['id']}/abrir", headers=manager)
    # Se retira el azufre del maestro
    client.delete(f"/api/articulos/{azufre}", headers=admin)

    nueva = client.get("/api/solicitudes/mia", headers=lider).json()
    nsid = nueva["solicitud"]["id"]
    assert nsid != sid
    origenes = client.get(f"/api/solicitudes/{nsid}/clonables", headers=lider).json()
    vista = {v["nombre"]: v["estado"] for v in origenes[0]["vista"]}
    # el aceite tiene histórico (capturado al aprobar) -> LISTO; el azufre fue retirado -> NO_DISPONIBLE
    assert vista == {"Aceite agrícola": "LISTO", "Azufre mineral": "NO_DISPONIBLE"}
    r = client.post(f"/api/solicitudes/{nsid}/clonar", headers=lider, json={"origen_id": sid}).json()
    assert r["agregados"] == 1 and r["omitidos"] == 1


def test_cierre_cancela_borradores(client, lider, admin, manager):
    sid = client.get("/api/solicitudes/mia", headers=lider).json()["solicitud"]["id"]
    vig = client.get("/api/vigencias/activa", headers=admin).json()
    assert client.post(f"/api/vigencias/{vig['id']}/cerrar", headers=admin).status_code == 403
    r = client.post(f"/api/vigencias/{vig['id']}/cerrar", headers=manager).json()
    assert r["canceladas"] == 1
    assert client.get(f"/api/solicitudes/{sid}", headers=lider).json()["estado"] == "CANCELADA"


def test_importar_listado_maestro(client, admin):
    wb = Workbook()
    ws = wb.active
    ws.title = "LM-2027"
    ws.append(["N.", "Lote", "Código\nUNSPSC", "Producto", "Unidad SECOP ", "Descripción"])
    ws.append([1, "Agric", "12162003", "Aceite agrícola", "L", "Ficha actualizada"])
    ws.append([1, "Agric", 10171702, "Yodo agrícola", "gar", "Yodo"])
    ws.append([1, "Agric", "123", "Malo", "UN", ""])
    ws.append([99, "Nope", "12345678", "Lote inexistente", "UN", ""])
    buf = io.BytesIO()
    wb.save(buf)
    data = {"hoja": "LM-2027", "desactivar_ausentes": "true"}
    prev = client.post("/api/articulos/importar", headers=admin, data=data,
                       files={"archivo": ("lm.xlsx", buf.getvalue())}).json()
    assert prev["resumen"] == {"leidas": 4, "nuevos": 1, "actualizados": 1, "sin_cambios": 0, "errores": 2,
                               "advertencias": 0, "a_desactivar": 3}
    r = client.post("/api/articulos/importar", headers=admin, data={**data, "confirmar": "true"},
                    files={"archivo": ("lm.xlsx", buf.getvalue())})
    assert r.json()["aplicado"]
    nombres = {a["nombre"]: a for a in client.get("/api/articulos", headers=admin).json()["items"]}
    assert set(nombres) == {"Aceite agrícola", "Yodo agrícola"}
    assert nombres["Yodo agrícola"]["unidad"] == "GAR"


def test_propuesta_articulo(client, lider, admin):
    p = client.post("/api/propuestas", headers=lider, json={"nombre": "Semilla de cilantro", "unidad": "KG",
                                                            "descripcion": "Semilla certificada"}).json()
    assert client.post(f"/api/propuestas/{p['id']}/responder", headers=admin,
                       json={"aprobar": True, "codigo_unspsc": "1015"}).status_code == 422
    r = client.post(f"/api/propuestas/{p['id']}/responder", headers=admin,
                    json={"aprobar": True, "codigo_unspsc": "10151500"}).json()
    assert r["estado"] == "APROBADA" and r["articulo_id"]
    assert "Semilla de cilantro" in _ids(client, lider)


def test_lider_no_ve_otros_lotes(client, admin):
    otro = login(client, "lider.viver@sena.edu.co", "Lider123!")
    assert client.get("/api/articulos", headers=otro).json()["total"] == 0
    aid = client.get("/api/articulos", headers=admin).json()["items"][0]["id"]
    assert client.get(f"/api/articulos/{aid}", headers=otro).status_code == 403
